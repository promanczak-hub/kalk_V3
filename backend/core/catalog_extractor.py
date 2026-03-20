"""Catalog variant extractor.

Extracts structured variant data from catalog PDFs (via Gemini Pro)
and price list XLSX files (via openpyxl).

Each variant in a catalog represents a specific vehicle configuration
(e.g., Crafter L3H3 FWD 177KM) with its features and specs.
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any

import openpyxl
from google.genai import types

from core.database import supabase as sb_client
from core.gemini_client import (
    SAFETY_SETTINGS_PERMISSIVE,
    get_gemini_client,
)
from core.json_utils import clean_json_response

logger = logging.getLogger(__name__)

_STORAGE_BUCKET = "catalog-documents"

# ── Catalog extraction prompt ────────────────────────────────────

_CATALOG_EXTRACTION_PROMPT = """Jesteś ekspertem od analizy katalogu pojazdów.
Twoim zadaniem jest wyciągnięcie WSZYSTKICH wariantów pojazdu z tego dokumentu.

Dla KAŻDEGO wariantu wyciągnij:
- variant_name: pełna nazwa wariantu (np. "Crafter 35 Furgon L3H3 FWD 177KM")
- body_type: typ nadwozia/zabudowy (Furgon, Skrzynia, Platforma, Kombi, etc.)
- length_class: klasa długości (L2, L3, L4, etc.) jeśli dotyczy
- height_class: klasa wysokości (H2, H3, etc.) jeśli dotyczy
- drive_type: rodzaj napędu (FWD, RWD, 4MOTION, AWD, etc.)
- engine_power_hp: moc silnika w KM
- engine_capacity: pojemność silnika
- fuel_type: rodzaj paliwa (Diesel, Benzyna, Elektryczny, etc.)
- gvw_kg: DMC (dopuszczalna masa całkowita) w kg
- payload_kg: ładowność w kg
- cargo_length_mm: długość przestrzeni ładunkowej w mm
- cargo_width_mm: szerokość przestrzeni ładunkowej w mm
- cargo_height_mm: wysokość przestrzeni ładunkowej w mm
- cargo_volume_m3: kubatura przestrzeni ładunkowej w m3
- europallets: ilość europalet (jeśli podana)
- cargo_area_m2: powierzchnia ładunkowa w m2 (jeśli podana)
- wheelbase_mm: rozstaw osi w mm
- overall_length_mm: długość całkowita pojazdu w mm
- overall_width_mm: szerokość całkowita pojazdu w mm
- overall_height_mm: wysokość całkowita pojazdu w mm
- standard_equipment: lista wyposażenia standardowego (nazwy)
- price_net: cena netto PLN (jeśli podana)
- price_gross: cena brutto PLN (jeśli podana)

Jeśli informacja nie jest dostępna w dokumencie, wstaw null.
Zwróć WSZYSTKIE warianty, nawet jeśli różnią się tylko silnikiem lub napędem. Oczekujemy podejścia "Best-Effort": dokument może być materiałem promocyjnym, broszurą a nie pełnym cennikiem. Nawet jeśli masz tylko 1 wariant z ograniczonymi danymi technicznymi, spróbuj go wyekstrahować.

WAŻNE:
- NIE pomijaj wariantów
- NIE łącz wariantów które różnią się napędem (FWD vs 4MOTION = 2 warianty)
- NIE łącz wariantów które różnią się długością/wysokością (L3H3 vs L4H2 = 2 warianty)
- Wymiary ładunkowe przypisuj do KONKRETNEGO wariantu (nie ogólnie)
- Jeśli europalety dotyczą danego wariantu, przypisz je do niego

Zwróć JSON:
{
  "brand": "Volkswagen",
  "model_family": "Crafter",
  "year": 2026,
  "variants": [
    {
      "variant_name": "...",
      "body_type": "...",
      ...
    }
  ]
}
"""


def _extract_pdf_catalog(
    file_bytes: bytes,
    catalog_meta: dict[str, Any],
) -> dict[str, Any]:
    """Extract variants from a PDF catalog using Gemini Pro.

    Uses the same model and config as the digital twin extraction
    but with a catalog-specific prompt.
    """
    client = get_gemini_client()

    contents = [
        types.Part.from_bytes(data=file_bytes, mime_type="application/pdf"),
        _CATALOG_EXTRACTION_PROMPT,
    ]

    config = types.GenerateContentConfig(
        temperature=0.0,
        seed=42,
        max_output_tokens=65536,
        response_mime_type="application/json",
        system_instruction=(
            "Jesteś parserem dokumentów motoryzacyjnych (cenników, katalogów, broszur). "
            "Wyciągasz strukturyzowane dane o wariantach, nawet fragmentowane. "
            "Odpowiadasz wyłącznie w formacie JSON."
        ),
        safety_settings=SAFETY_SETTINGS_PERMISSIVE,
        thinking_config=types.ThinkingConfig(thinking_budget=16384),
    )

    logger.info(
        "Extracting catalog variants via Gemini Pro: %s",
        catalog_meta.get("display_name", "unknown"),
    )

    response = client.models.generate_content(
        model="gemini-2.5-pro",
        contents=contents,
        config=config,
    )

    raw_text = response.text or ""
    cleaned = clean_json_response(raw_text)
    data = json.loads(cleaned)

    variants = data.get("variants", [])
    logger.info("Extracted %d variants from PDF catalog", len(variants))

    return {
        "extracted_data": data,
        "variant_count": len(variants),
    }


def _extract_xlsx_catalog(
    file_bytes: bytes,
    catalog_meta: dict[str, Any],
) -> dict[str, Any]:
    """Extract variants from an XLSX price list using openpyxl.

    Heuristic: looks for sheets with vehicle model data
    (columns with brand, model, price, variant info).
    """
    wb = openpyxl.load_workbook(
        io.BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    all_variants: list[dict[str, Any]] = []

    for ws in wb.worksheets:
        # Try to find header row (look in first 5 rows)
        header_row: list[str] = []
        header_idx = 0

        for row_idx in range(1, 6):
            row_vals = [
                str(cell.value).strip().lower()
                for cell in ws[row_idx]
                if cell.value is not None
            ]
            # Heuristic: header has keywords like model, cena, wariant
            keywords = {"model", "wariant", "cena", "moc", "silnik", "typ"}
            matches = sum(1 for v in row_vals if any(k in v for k in keywords))
            if matches >= 2:
                header_row = [
                    str(c.value).strip() if c.value else f"col_{i}"
                    for i, c in enumerate(ws[row_idx])
                ]
                header_idx = row_idx
                break

        if not header_row or not header_idx:
            continue

        # Read data rows
        for row in ws.iter_rows(
            min_row=header_idx + 1,
            max_row=ws.max_row,
            values_only=False,
        ):
            row_dict: dict[str, Any] = {}
            has_data = False
            for i, cell in enumerate(row):
                if i < len(header_row) and cell.value is not None:
                    row_dict[header_row[i]] = cell.value
                    has_data = True
            if has_data:
                all_variants.append(
                    {
                        "variant_name": _build_variant_name(row_dict),
                        "raw_data": row_dict,
                        "source_sheet": ws.title,
                    }
                )

    wb.close()

    logger.info(
        "Extracted %d rows from XLSX: %s",
        len(all_variants),
        catalog_meta.get("display_name", "unknown"),
    )

    return {
        "extracted_data": {
            "brand": catalog_meta.get("brand", ""),
            "model_family": catalog_meta.get("model_family", ""),
            "variants": all_variants,
        },
        "variant_count": len(all_variants),
    }


def _build_variant_name(row: dict[str, Any]) -> str:
    """Build a variant name from row data using common column names."""
    parts: list[str] = []
    name_keys = ["model", "wariant", "wersja", "nazwa", "opis"]
    for key in name_keys:
        for col_name, val in row.items():
            if key in col_name.lower() and val:
                parts.append(str(val).strip())
                break

    return " ".join(parts) if parts else str(list(row.values())[:3])


def _extract_pricelist_hybrid(
    file_bytes: bytes,
    markdown_content: str,
    catalog_meta: dict[str, Any],
) -> dict[str, Any]:
    """Extract variants using the hybrid PricingAgent for Price Lists."""
    logger.info(
        "Extracting price list via PricingAgent (Hybrid): %s",
        catalog_meta.get("display_name", "unknown"),
    )

    from core.pdf_pipeline.agents import PricingAgent

    agent = PricingAgent()

    parsed = agent.extract_data(markdown_content, file_bytes)

    variants = []
    for engine in parsed.engines:
        for trim in engine.prices_by_trim:
            if trim.price_netto or trim.price_brutto:
                variants.append(
                    {
                        "variant_name": f"{engine.engine_name} {trim.trim_name}".strip(),
                        "body_type": None,
                        "engine_power_hp": engine.power_hp,
                        "fuel_type": engine.fuel_type,
                        "price_net": trim.price_netto,
                        "price_gross": trim.price_brutto,
                        "transmission": engine.transmission,
                        "standard_equipment": [],
                    }
                )

    logger.info("Mapped PricingAgent output to %d flat variants", len(variants))

    return {
        "extracted_data": {
            "brand": parsed.brand or catalog_meta.get("brand", ""),
            "model_family": parsed.model or catalog_meta.get("model_family", ""),
            "year": parsed.model_year,
            "variants": variants,
        },
        "variant_count": len(variants),
    }


# ── Main entry point ─────────────────────────────────────────────


def extract_catalog_variants(
    catalog: dict[str, Any],
) -> dict[str, Any]:
    """Extract variants from a catalog document.

    Dispatches to PDF or XLSX extractor based on file_type.
    Downloads user file from Supabase Storage.

    Args:
        catalog: Full catalog row from model_document_sources.

    Returns:
        Dict with extracted_data and variant_count.
    """
    storage_path = catalog["storage_path"]
    file_type = catalog["file_type"]
    document_type = catalog.get("document_type")

    # Download file from storage
    file_bytes = sb_client.storage.from_(_STORAGE_BUCKET).download(storage_path)

    if file_type == "pdf":
        if document_type == "price_list":
            markdown_content = catalog.get("document_markdown", "")
            if not markdown_content:
                logger.warning(
                    "No markdown found for hybrid processing, passing empty string"
                )
            return _extract_pricelist_hybrid(
                file_bytes, markdown_content or "", catalog
            )
        return _extract_pdf_catalog(file_bytes, catalog)
    if file_type in ("xlsx", "csv"):
        return _extract_xlsx_catalog(file_bytes, catalog)

    raise ValueError(f"Unsupported file type for extraction: {file_type}")
