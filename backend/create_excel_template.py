import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import os


def create_template():
    # Ensure templates directory exists
    os.makedirs(r"d:\kalk_v3\backend\templates", exist_ok=True)
    file_path = r"d:\kalk_v3\backend\templates\template_oferta.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zestawienie"

    # Set up some headers
    headers = [
        "Marka",
        "Model",
        "Silnik/Powertrain",
        "VIN / Kod Konfiguracji",
        "Okres (m-ce)",
        "Przebieg (km/rok)",
        "Rata Netto (PLN)",
        "Rekomendacja Systemu",
    ]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003366")
        cell.alignment = Alignment(horizontal="center", vertical="center")

        # Adjust some specific column widths
        col_letter = openpyxl.utils.get_column_letter(col_num)
        if header in ["Silnik/Powertrain", "VIN / Kod Konfiguracji"]:
            ws.column_dimensions[col_letter].width = 30
        else:
            ws.column_dimensions[col_letter].width = 20

    # Title
    ws["A1"] = "OFERTA FLOTOWA"
    ws["A1"].font = Font(size=20, bold=True)
    ws["A2"] = "Dla Klienta:"
    ws["B2"] = "[NAZWA_KLIENTA]"
    ws["A3"] = "NIP:"
    ws["B3"] = "[NIP_KLIENTA]"

    wb.save(file_path)
    print(f"Szablon został zapisany w {file_path}")


if __name__ == "__main__":
    create_template()
