import os
from openpyxl import load_workbook

def create_real_offer(template_path: str, out_path: str):
    wb = load_workbook(template_path)
    ws = wb.active
    
    # 3 przykładowe auta
    cars = [
        {
            "code": "KALK-2026-001",
            "model_str": "Toyota Corolla 1.8 Hybrid",
            "options": "Pakiet Comfort, Lakier metalizowany",
            "term": "36 mc",
            "mileage": 20000,
            "total_price": 1500.0,
            "fin_price": 1200.0,
            "tech_price": 300.0,
            "overmileage": 0.45,
            "cost_type": "Pełny Serwis"
        },
        {
            "code": "KALK-2026-002",
            "model_str": "Skoda Octavia 2.0 TDI",
            "options": "Pakiet Style",
            "term": "48 mc",
            "mileage": 30000,
            "total_price": 1800.0,
            "fin_price": 1400.0,
            "tech_price": 400.0,
            "overmileage": 0.50,
            "cost_type": "Pełny Serwis"
        },
        {
            "code": "KALK-2026-003",
            "model_str": "Porsche Macan 2.0 TSI",
            "options": "Felgi 20', Zawieszenie PASM",
            "term": "24 mc",
            "mileage": 15000,
            "total_price": 4500.0,
            "fin_price": 3900.0,
            "tech_price": 600.0,
            "overmileage": 1.50,
            "cost_type": "Pełny Serwis"
        }
    ]
    
    start_row = 12
    
    for i, car in enumerate(cars):
        row = start_row + i
        ws.cell(row=row, column=1, value=car["code"]) # Kod kalkulacji
        ws.cell(row=row, column=2, value=car["model_str"]) # Model samochodu
        ws.cell(row=row, column=3, value=car["options"]) # Wyposażenie fabryczne
        ws.cell(row=row, column=4, value=car["term"]) # Okres umowy
        ws.cell(row=row, column=5, value=car["mileage"]) # Limit km
        
        # Wpisujemy formułę, żeby upewnić się, że Excel ładnie zsumuje ratę
        ws.cell(row=row, column=6, value=f"=G{row}+H{row}") # Miesięczna cena (Finansowa + Techniczna)

        ws.cell(row=row, column=7, value=car["fin_price"]) # Część finansowa
        ws.cell(row=row, column=8, value=car["tech_price"]) # Część techniczna
        ws.cell(row=row, column=9, value=car["overmileage"]) # Opłata za nadprzebieg
        ws.cell(row=row, column=10, value=car["cost_type"]) # Rodzaj kosztów
        
    wb.save(out_path)
    print(f"Pomyślnie utworzono {out_path}")

if __name__ == "__main__":
    template = os.path.join(os.getcwd(), "templates", "extracted", "backend", "app", "core", "template_v6.xlsx")
    out = os.path.join(os.getcwd(), "Gotowa_Oferta_Testowa.xlsx")
    create_real_offer(template, out)
