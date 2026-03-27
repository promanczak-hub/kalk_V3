import openpyxl

path = r"C:\Users\proma\Downloads\DRAFT_KALKULATORA_WARTOŚCI_REZYDUALNYCH_ver_aktualna_JŁ_02_02_2026 (3).xlsx"
try:
    wb = openpyxl.load_workbook(path, data_only=False)
except Exception as e:
    print("Failed to load:", e)
    exit(1)

sheet_name = "KALKULATOR DH (dubel)"
if sheet_name not in wb.sheetnames:
    print(f"Sheet not found. Available: {wb.sheetnames}")
    exit(1)

ws = wb[sheet_name]

cols = [
    "A",
    "B",
    "P",
    "Q",
    "R",
    "S",
    "T",
    "W",
    "X",
    "Y",
    "Z",
    "AA",
    "AB",
    "AC",
    "AD",
    "AE",
    "AI",
    "AJ",
    "AK",
    "AL",
    "AM",
    "AN",
    "AO",
    "AP",
    "AQ",
]

print("FORMULAS (data_only=False):")
for col in cols:
    header = ws[f"{col}1"].value
    val = ws[f"{col}1679"].value
    print(f"{col}1679 [{header}]: {val}")

print("\nVALUES (data_only=True):")
wb2 = openpyxl.load_workbook(path, data_only=True)
ws2 = wb2[sheet_name]
for col in cols:
    header = ws2[f"{col}1"].value
    val = ws2[f"{col}1679"].value
    print(f"{col}1679 [{header}]: {val}")
