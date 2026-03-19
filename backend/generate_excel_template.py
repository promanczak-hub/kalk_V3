import os
from openpyxl import load_workbook

def create_offers_from_template(template_path: str, out_path: str):
    # Wczytaj szablon
    wb = load_workbook(template_path)
    ws = wb.active
    
    # Przykładowe dane 3 aut dopasowane do kolumn zdefiniowanych w szablonie
    # Z poprzedniego odczytu wiemy, że nagłówki są w wierszu 5 (A5:H5)
    # Piszemy od wiersza 6.
    
    cars = [
        {
            "brand": "Toyota",
            "model": "Corolla",
            "powertrain": "1.8 Hybrid",
            "vin": "BRAK",
            "term": 36,
            "mileage": 20000,
            "installment": 1500.0,
            "recommendation": "Szczególnie polecane"
        },
        {
            "brand": "Skoda",
            "model": "Octavia",
            "powertrain": "2.0 TDI",
            "vin": "BRAK",
            "term": 48,
            "mileage": 30000,
            "installment": 1800.0,
            "recommendation": "Dobry wybór flotowy"
        },
        {
            "brand": "Porsche",
            "model": "Macan",
            "powertrain": "2.0 TSI",
            "vin": "BRAK",
            "term": 24,
            "mileage": 15000,
            "installment": 4500.0,
            "recommendation": "Premium"
        }
    ]
    
    start_row = 6
    
    for i, car in enumerate(cars):
        row = start_row + i
        ws.cell(row=row, column=1, value=car["brand"])
        ws.cell(row=row, column=2, value=car["model"])
        ws.cell(row=row, column=3, value=car["powertrain"])
        ws.cell(row=row, column=4, value=car["vin"])
        ws.cell(row=row, column=5, value=car["term"])
        ws.cell(row=row, column=6, value=car["mileage"])
        ws.cell(row=row, column=7, value=car["installment"])
        ws.cell(row=row, column=8, value=car["recommendation"])
        
    wb.save(out_path)
    print(f"Zapisano plik Excel: {out_path} na bazie szablonu {template_path}")

if __name__ == "__main__":
    template_path = os.path.join(os.getcwd(), "templates", "template_oferta.xlsx")
    out_path = os.path.join(os.getcwd(), "oferty_koszyk_wzorcowy.xlsx")
    create_offers_from_template(template_path, out_path)
