import openpyxl

def main():
    file_path = r"C:\Users\proma\Downloads\DRAFT_KALKULATORA_WARTOŚCI_REZYDUALNYCH_ver_aktualna_JŁ_02_02_2026_final.xlsx"
    sheet_name = "KALKULATOR DH (dubel)"
    target_row = 1682

    out_file = r"d:\kalk_v3\backend\scripts\excel_output_utf8.txt"

    with open(out_file, "w", encoding="utf-8") as f:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=False)
            
            if sheet_name not in wb.sheetnames:
                f.write(f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}\n")
                return
                
            sheet = wb[sheet_name]
            
            wb_data = openpyxl.load_workbook(file_path, data_only=True)
            sheet_data = wb_data[sheet_name]
            
            f.write(f"Reading row {target_row} from sheet '{sheet_name}'...\n")
            
            headers = []
            for cell in sheet[1]:
                headers.append(cell.value)
                
            f.write("-" * 50 + "\n")
            for col_idx in range(1, sheet.max_column + 1):
                cell_formula = sheet.cell(row=target_row, column=col_idx)
                cell_data = sheet_data.cell(row=target_row, column=col_idx)
                
                if cell_formula.value is None and cell_data.value is None:
                    continue
                    
                col_letter = cell_formula.column_letter
                header = headers[col_idx - 1] if col_idx - 1 < len(headers) else f"Col {col_letter}"
                
                formula_text = str(cell_formula.value) if cell_formula.value is not None else ""
                data_text = str(cell_data.value) if cell_data.value is not None else ""
                
                f.write(f"[{col_letter}] {header}:\n")
                f.write(f"  Formula: {formula_text}\n")
                if formula_text != data_text:
                   f.write(f"  Value:   {data_text}\n")
                f.write("\n")
                
        except Exception as e:
            f.write(f"Error reading excel: {e}\n")

if __name__ == "__main__":
    main()
