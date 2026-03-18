import openpyxl
from openpyxl.utils import get_column_letter

file_path = r"C:\Users\proma\Downloads\DRAFT_KALKULATORA_WARTOŚCI_REZYDUALNYCH_ver_aktualna_JŁ_02_02_2026_final.xlsx"

try:
    wb = openpyxl.load_workbook(file_path, data_only=False)
    ws = wb["KALKULATOR DH (dubel)"]
    print("--- ROW 1682 FORMULAS ---")
    
    row_1682 = [cell.value for cell in ws[1682]]
    
    cols_of_interest = ["L", "O", "P", "AB", "AD", "AF", "AI", "AJ", "AK", "AN", "AR", "AV", "AZ", "BC", "BG", "BK", "BO", "BQ", "BR", "BS", "BT", "BU", "BV", "BX", "BY"]
    
    with open("row_1682_formulas.txt", "w", encoding="utf-8") as f:
        for col_idx, val in enumerate(row_1682):
            col_letter = get_column_letter(col_idx + 1)
            if col_letter in cols_of_interest:
                val_str = str(val).replace('\r', '').replace('\n', ' ')
                if val_str.startswith("="):
                    f.write(f"[{col_letter}] = {val_str}\n")
                else:
                    f.write(f"[{col_letter}] RAW: {val_str}\n")
    print("Wrote formulas to row_1682_formulas.txt")
except Exception as e:
    print("Error:", e)
