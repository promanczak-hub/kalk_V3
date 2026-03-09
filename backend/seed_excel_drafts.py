import uuid
import openpyxl
from core.database import supabase

excel_path = r"C:\Users\proma\Downloads\DRAFT KALKULATORA WARTOŚCI REZYDUALNYCH ver aktualna JŁ 02.02 (version 1).xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 8 middle sheets
sheets_to_process = wb.sheetnames[1:-1]

total_saved = 0
for sheet_name in sheets_to_process:
    ws = wb[sheet_name]

    # We will treat the first row as headers.
    headers = []
    max_col = ws.max_column
    for c in range(1, max_col + 1):
        val = ws.cell(1, c).value
        # If header is None or empty string, name it "Column X"
        hdr_str = str(val).strip() if val is not None else f"Col_{c}"
        if not hdr_str:
            hdr_str = f"Col_{c}"
        headers.append({"field": f"col_{c}", "headerName": hdr_str, "width": 150})

    data_rows = []
    # Start from row 2 for data
    for r in range(2, ws.max_row + 1):
        row_obj = {"id": str(uuid.uuid4())}
        is_empty = True
        for c in range(1, max_col + 1):
            val = ws.cell(r, c).value
            field_name = f"col_{c}"
            row_obj[field_name] = val if val is not None else ""
            if val is not None and str(val).strip() != "":
                is_empty = False

        # Only add row if it's not completely empty
        if not is_empty:
            data_rows.append(row_obj)

    print(f"Sheet '{sheet_name}': {len(headers)} columns, {len(data_rows)} rows")

    # Insert or update in Supabase
    record = {"sheet_name": sheet_name, "columns_def": headers, "data_rows": data_rows}

    # Upsert using sheet_name as unique key
    try:
        response = (
            supabase.table("excel_drafts")
            .upsert(record, on_conflict="sheet_name")
            .execute()
        )
        print("  -> Saved to Supabase")
        total_saved += 1
    except Exception as e:
        print(f"  -> Error saving to Supabase: {e}")

print(f"Done parsing and uploading {total_saved} Excel drafts.")
