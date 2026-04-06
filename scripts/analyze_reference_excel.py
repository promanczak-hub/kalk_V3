import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import json

file_path = r'C:\Users\proma\Downloads\Oferta_2026-04-01_ROBELIT_PRO_SPÓŁKA_Z_OGRANICZONĄ_ODPOWIEDZIALNOŚCIĄ.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb.active

def get_color(color):
    if color is None:
        return None
    try:
        return str(color.rgb)
    except:
        return str(getattr(color, 'theme', None))

data = []
for row in range(1, min(30, ws.max_row + 1)):
    row_data = []
    for col in range(1, min(15, ws.max_column + 1)):
        cell = ws.cell(row=row, column=col)
        
        font_info = None
        if cell.font:
            font_info = {
                'name': cell.font.name,
                'size': cell.font.size,
                'bold': cell.font.bold,
                'italic': cell.font.italic,
                'color': get_color(cell.font.color) 
            }
            
        fill_info = None
        if cell.fill and cell.fill.start_color:
            fill_info = get_color(cell.fill.start_color)
            
        alignment_info = None
        if cell.alignment:
            alignment_info = {
                'horizontal': cell.alignment.horizontal,
                'vertical': cell.alignment.vertical
            }
            
        cell_info = {
            'addr': cell.coordinate,
            'val': str(cell.value) if cell.value is not None else None,
            'font': font_info,
            'fill': fill_info,
            'align': alignment_info
        }
        row_data.append(cell_info)
    data.append(row_data)

widths = {openpyxl.utils.get_column_letter(i): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width for i in range(1, 15)}

output = {
    'widths': widths,
    'rows': data
}

with open('excel_style_dump.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, indent=2, ensure_ascii=False)
print("Dumped to excel_style_dump.json")
