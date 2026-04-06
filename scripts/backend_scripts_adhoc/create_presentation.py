import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

# Slide 1: Title
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Kalkulator V3 vs Architektura V1"
subtitle.text = "Porównanie systemu, bezpieczeństwa marży i architektury kalkulacyjnej\nPrezentacja dla Zarządu"

# Slide 2: Główne Różnice - Dlaczego wybudowaliśmy V3?
bullet_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Biznesowe Cele Transformacji (Dlaczego V3?)"
tf = body_shape.text_frame
tf.text = "Bezpieczeństwo finansowe i Data-Driven:"
p = tf.add_paragraph()
p.text = "Zero Hardkodowania: Odeszliśmy od wpisywania na stałe konkretnych marek/modeli i kwot w kodzie źródłowym."
p.level = 1
p = tf.add_paragraph()
p.text = "Elastyczność: Wszystkie parametry i cenniki pobierane są dynamicznie z dedykowanej bazy danych."
p.level = 1

p = tf.add_paragraph()
p.text = "Ochrona przed 'Cichymi Błędami' (Zasada Fail-Fast):"
p = tf.add_paragraph()
p.text = "W V1 system często kontynuował obliczenia korzystając z domyślnych stałych gdy brakowało danych."
p.level = 1
p = tf.add_paragraph()
p.text = "W V3, w razie braku stawki (np. ubezpieczenia), wyliczenie jest natychmiast zatrzymywane z jasnym alarmem, ucinając ryzyko błędu w marży."
p.level = 1

p = tf.add_paragraph()
p.text = "Pełna Transparentność (Ślad Rewizyjny):"
p = tf.add_paragraph()
p.text = "Każda wyliczona liczba na dokumencie końcowym ma swój cyfrowy dowód matematyczny (z jakiego działania się dokładnie wzięła)."
p.level = 1

# Slide 3: Tabela: Orkiestrator i część subkalkulatorów
blank_slide_layout = prs.slide_layouts[5]  # Title Only
slide = prs.slides.add_slide(blank_slide_layout)
shapes = slide.shapes
shapes.title.text = "Różnice w działaniu: Orkiestrator i Kalkulatory (1/2)"

rows = 6
cols = 3
left = Inches(0.2)
top = Inches(1.3)
width = Inches(9.6)
height = Inches(5.0)

table = shapes.add_table(rows, cols, left, top, width, height).table

table.columns[0].width = Inches(2.0)
table.columns[1].width = Inches(3.8)
table.columns[2].width = Inches(3.8)

headers = [
    "Komponent",
    "Kalkulator V1 (Mechanizm Zastany)",
    "Kalkulator V3 (Nowoczesny)",
]
for col_idx, header in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.text = header
    cell.text_frame.paragraphs[0].font.bold = True
    cell.text_frame.paragraphs[0].font.size = Pt(13)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0x00, 0x33, 0x66)
    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

data_sl1 = [
    (
        "Orkiestrator (Główny silnik)",
        "Ciężki organizm współdzielący ten sam stan dla wszystkich modułów. Podatny na niezauważone nadpisywanie zmiennych.",
        "Restrykcyjny Pipeline (12 ściśle oddzielonych kroków). Każdy subkalkulator to szczelna fabryka posiadająca wejście i konkretne wyjście. Jawny wstępny audyt danych (Readiness Check).",
    ),
    (
        "Sub: Opony",
        "Składniki hardkodowane na sztywno, ukryte progi i domniemane ceny za sztukę.",
        "Sterowany elastyczną bazą danych. Obsługa 11 rozmiarów x 13 kategorii opon i konkretnych progów przebiegowych odcinających zużycie. Kalkulacja za komplet, gotowa operacyjnie.",
    ),
    (
        "Sub: Koszty Dodatkowe",
        "Wszelkie dotacje i ryczałty zaszyte w logice kodu.",
        "Uporządkowany izolowany moduł, uwzględnia precyzyjnie przygotowanie pojazdu i łatwość manipulacji kosztem czynszu administracyjnego przed wrzuceniem do CAPEXu.",
    ),
    (
        "Sub: Samochód Zastępczy",
        "Rozmyta odpowiedzialność w obliczeniach zintegrowanych z innymi serwisami.",
        "Wysłany do osobnego kontenera obliczeniowego, parametry i wirtualne stawki potwierdzone wprost, bardzo ułatwione debugowanie biznesowe.",
    ),
    (
        "Sub: Serwis",
        "Ukryte progi miesięczne (km), w przypadku błędów milczał.",
        "Strukturalna analiza wg przebiegu klasowego (minimalny próg 1667 km/mc twardo wyegzekwowany). Korekta kosztów (% power-bandu) wydzielona i w pełni audytowalna.",
    ),
]

for row_idx, row_data in enumerate(data_sl1, start=1):
    for col_idx, cell_value in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        cell.text = cell_value
        cell.text_frame.paragraphs[0].font.size = Pt(11)


# Slide 4: Tabela: Pozostałe subkalkulatory
slide = prs.slides.add_slide(blank_slide_layout)
shapes = slide.shapes
shapes.title.text = "Różnice w działaniu: Pozostałe Kalkulatory (2/2)"

rows = 6
cols = 3
table2 = shapes.add_table(rows, cols, left, top, width, height).table

table2.columns[0].width = Inches(2.0)
table2.columns[1].width = Inches(3.8)
table2.columns[2].width = Inches(3.8)

for col_idx, header in enumerate(headers):
    cell = table2.cell(0, col_idx)
    cell.text = header
    cell.text_frame.paragraphs[0].font.bold = True
    cell.text_frame.paragraphs[0].font.size = Pt(13)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0x00, 0x33, 0x66)
    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

data_sl2 = [
    (
        "Sub: Ubezpieczenie",
        "Pętla 7-letnia, 'ciche ratowanie' zniżek; gdy system nie pobrał stawki AC/OC dla 5-roku, to generował zero, psując marżę całej umowy.",
        "Zatrzymana Pętla 7-letnia. Całkowity zakaz 'Fallbacków'. System odmawia przeliczenia, ratując biznes przed sprzedażą i raportując dokładny brak, np. 'Brak stawki dla Klasy Premium rocznik 2025'.",
    ),
    (
        "Sub: Utrata Wartości (RV)",
        "Bardzo złożony, mętny black-box pełen instrukcji opartych o stare ID aut. Brak jasności jak wyliczona kwota powstała.",
        "V1 Parity: Matematyka przełożona 1:1 ze starego sytemu dla bezpieczeństwa, jednak z nałożonym 'Śladem Rewizyjnym'. Widzimy każdą wziętą do potrącenia oponę z osobna na czytelnym wyciągu JSON.",
    ),
    (
        "Sub: Cena Zakupu (CAPEX)",
        "Brak struktury wejścia ceny netto vs transport. Rozliczenia podatkowe mocno chaotyczne.",
        "Twardy fundament netto-based CAPEX. Czysty podział na pozycje niepodlegające rabatowaniu (non-discountable) i podstawę rabatową, ułatwia dogadanie z dilerami.",
    ),
    (
        "Sub: Finanse & Amortyzacja",
        "Algorytm PMT z gąszczem ifów blokujących różne typy wkładów.",
        "Elastyczne i w pełni izolowane od czynników zewnętrznych wzory matematyczne, odporne na podanie okresów ujemnych.",
    ),
    (
        "Sub: Budżet Mktg & Koszt Dz.",
        "Sklejone z całością raportowania, nieistniejące w postaci autonomicznych wyliczeń.",
        "Samodzielne mikroserwisy (jedno przejrzyste mnożenie WR x VAT x budżet %), dające się łatwo modelować w przyszłości bez psucia głównego trzonu kalkulacji.",
    ),
]

for row_idx, row_data in enumerate(data_sl2, start=1):
    for col_idx, cell_value in enumerate(row_data):
        cell = table2.cell(row_idx, col_idx)
        cell.text = cell_value
        cell.text_frame.paragraphs[0].font.size = Pt(11)

# Slide 5: Konkluzje dla Zarządu
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Podsumowanie dla Biznesu"
tf = body_shape.text_frame
tf.text = "Pełna kontrola nad rentownością:"
p = tf.add_paragraph()
p.text = "Koniec z ukrytymi stratami wynikającymi z nieaktualnych stawek. System w V3 raczej odmówi wykonania operacji, niż wypuści ofertę opartą o błędne lub darmowe (zero) parametry serwisowe/AC."
p.level = 1

p = tf.add_paragraph()
p.text = "Agility (Szybkość Dostosowania):"
p = tf.add_paragraph()
p.text = "Nasz zespół może konfigurować nowe produkty prosto przez dynamiczny Panel Sterowania bez zlecania tygodni pracy programistycznej."
p.level = 1

p = tf.add_paragraph()
p.text = "Solidny punkt wyjścia dla AI i Automatyzacji:"
p = tf.add_paragraph()
p.text = "Czysta, ścisła architektura pozwala na szybkie zapięcie zaawansowanych algorytmów analizy rynku (np. auto-czytanie cenników PDF za pomocą Vertex/Docling), bez narażania precyzji samego serca finansowego."
p.level = 1

output_file = os.path.join("d:\\kalk_v3", "Prezentacja_Kalkulator_V1_vs_V3.pptx")
prs.save(output_file)
print(f"Prezentacja zapisana w: {output_file}")
