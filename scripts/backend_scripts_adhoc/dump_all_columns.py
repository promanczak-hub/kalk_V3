import openpyxl

path = r"C:\Users\proma\Downloads\DRAFT_KALKULATORA_WARTOŚCI_REZYDUALNYCH_ver_aktualna_JŁ_02_02_2026 (3).xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["KALKULATOR DH (dubel)"]

vals = {}
for col_idx in range(1, 100):
    val_header = ws.cell(row=1, column=col_idx).value
    val_1679 = ws.cell(row=1679, column=col_idx).value
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    vals[col_letter] = (val_header, val_1679)

with open("dump_excel.txt", "w", encoding="utf-8") as f:
    for k, v in vals.items():
        f.write(f"{k}1679 [{v[0]}]: {v[1]}\n")

wb_formulas = openpyxl.load_workbook(path, data_only=False)
ws_formulas = wb_formulas["KALKULATOR DH (dubel)"]
with open("dump_excel_formulas.txt", "w", encoding="utf-8") as f:
    for col_idx in range(1, 100):
        val_header = ws_formulas.cell(row=1, column=col_idx).value
        val_1679 = ws_formulas.cell(row=1679, column=col_idx).value
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        f.write(f"{col_letter}1679 [{val_header}]: {val_1679}\n")
