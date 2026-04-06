import openpyxl

file_path = r'C:\Users\proma\Downloads\Oferta_2026-04-01_ROBELIT_PRO_SPÓŁKA_Z_OGRANICZONĄ_ODPOWIEDZIALNOŚCIĄ.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=True)
ws = wb.active

with open('excel_clean_summary.txt', 'w', encoding='utf-8') as f:
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            if cell.value is not None:
                val = str(cell.value).strip()
                if val:
                    f.write(f"[{cell.coordinate}] {val}\n")
